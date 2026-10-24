"""Fill the recreated academic worksheets with an explicit planning scenario.

Only the green inventory contains company data. All financial inputs and
productivity forecasts below are editable analyst estimates, authorized by the
user, not payroll records, paid invoices, observed timings or system demo data.
"""
from pathlib import Path
import sys, json, math, hashlib, tempfile, zipfile
from datetime import datetime
sys.path.insert(0,str(Path(__file__).resolve().parent))
import build_reference_recreation as R

OUT=R.OUT;VAL=R.VAL;VIEWS=R.VIEWS
INPUTS={
 'dev1_rate':10000,'dev1_months':3,'dev2_rate':20000,'dev2_months':3,
 'dev3_rate':5000,'dev3_months':1,'dev4_rate':6000,'dev4_months':2,
 'training_rate':1500,'training_days':2,'training_kit':1000,
 'server_rate':1000,'server_qty':1,'drive_rate':2500,'drive_qty':1,
 'license_rate':0,'license_qty':1,
 'maintenance_rate':750,'frontend_rate':0,'communication_rate':200,'domain_rate':900,'renewal_rate':0,
 'fx_php_usd':65,'database_usd_month':25,'hourly_labor_rate':150,
 'existing_labor1_hours':40,'existing_labor2_hours':16,'existing_labor3_hours':24,
 'proposed_labor1_hours':20,'proposed_labor2_hours':6,'proposed_labor3_hours':12,
}
for prefix,values in {
 'existing':[(4,150),(3,150),(12,250),(1,150),(6,80),(12,60),(6,350)],
 'proposed':[(2,150),(2,150),(4,250),(2,80),(0,400),(3,350)]}.items():
    for i,(qty,price) in enumerate(values,1):
        INPUTS[f'{prefix}_supply{i}_qty']=qty;INPUTS[f'{prefix}_supply{i}_price']=price

SOURCES=[
 {'id':'S1','title':'Supabase pricing','url':'https://supabase.com/pricing','checked':'2026-10-24',
  'finding':'Pro starts at USD 25/month; first Micro project credit, database/file-storage quotas and seven-day daily database backups are included within published limits.',
  'use':'USD 25 x PHP 65/USD planning conversion = PHP 1,625/month. No separate database or bundled database-backup fee. Usage taxes/card charges/overages must be confirmed before purchase.'},
 {'id':'S2','title':'Cloudflare Pages limits','url':'https://developers.cloudflare.com/pages/platform/limits/','checked':'2026-10-24',
  'finding':'A Free plan exists, with published build and static-site limits.','use':'Static frontend hosting budget is PHP 0 within free static limits; backend service is budgeted separately.'},
 {'id':'S3','title':'Cloudflare Registrar','url':'https://www.cloudflare.com/domains/','checked':'2026-10-24',
  'finding':'Domain registration/renewal is at registry/ICANN cost; DNS and SSL are included.','use':'PHP 900/year is an analyst domain allowance, not a quoted domain price. SSL and open-source tool license budget is PHP 0.'},
 {'id':'S4','title':'NWPC NCR wage information','url':'https://nwpc.dole.gov.ph/ncr/','checked':'2026-10-24',
  'finding':'The page lists a PHP 755/day non-agriculture benchmark under NCR-28 effective 26 September 2026.',
  'use':'Only a Philippine costing reasonableness check. The PHP 150/hour administrative valuation is an analyst assumption, not a company wage or legal compliance conclusion. Project development retainers cover part-time services, not full-time salaries.'},
 {'id':'S5','title':'BSP exchange-rate statistics','url':'https://www.bsp.gov.ph/SitePages/Statistics/exchangerate.aspx','checked':'2026-10-24',
  'finding':'Official exchange-rate reference resource.','use':'PHP 65/USD is a rounded planning allowance, not a promised transaction rate. Replace it with the actual bank settlement rate when procuring hosting.'},
]

def expected():
    development=sum(INPUTS[f'dev{i}_rate']*INPUTS[f'dev{i}_months'] for i in range(1,5))
    training=INPUTS['training_rate']*INPUTS['training_days']+INPUTS['training_kit']
    hardware=sum(INPUTS[k+'_rate']*INPUTS[k+'_qty'] for k in ['server','drive','license'])
    recurring=12*(INPUTS['maintenance_rate']+INPUTS['database_usd_month']*INPUTS['fx_php_usd']+INPUTS['frontend_rate']+INPUTS['communication_rate'])+INPUTS['domain_rate']+INPUTS['renewal_rate']
    supplies={p:sum(INPUTS[f'{p}_supply{i}_qty']*INPUTS[f'{p}_supply{i}_price'] for i in range(1,n+1)) for p,n in [('existing',7),('proposed',6)]}
    labor={p:12*INPUTS['hourly_labor_rate']*sum(INPUTS[f'{p}_labor{i}_hours'] for i in range(1,4)) for p in ['existing','proposed']}
    cash=supplies['existing']-supplies['proposed'];productivity=labor['existing']-labor['proposed']
    return {'development_total':development,'training_total':training,'hardware_total':hardware,
      'one_time_total':development+training+hardware,'recurring_total':recurring,
      'existing_supplies_total':supplies['existing'],'proposed_supplies_total':supplies['proposed'],
      'existing_labor_total':labor['existing'],'proposed_labor_total':labor['proposed'],
      'supplies_reduction':cash,'labor_reduction':productivity,'total_reduction':cash+productivity,
      'annual_net_cash':cash-recurring,'annual_net_economic':cash+productivity-recurring}

EXPECTED=expected()

def alter_text(scene,old,new):
    matches=[b for b in scene.blocks if b.text==old]
    assert matches,(scene.name,old)
    for b in matches:b.text=new

def recurring(scene,x,y,duplicate=False):
    rows=[('Software Maintenance','maintenance',12),('Cloud Database + Storage','database',12),
          ('Static Frontend Hosting','frontend',12),('Incremental Communication','communication',12),
          ('Domain Renewal (year)','domain',1),('Software License Renewal','renewal',1)]
    prefix='copy_' if duplicate else ''
    R.table_row(scene,x,y,[117,58,53],19,['Description','Rate (PHP)','Annually'],True,R.GRAY,size=9)
    for i,(name,key,freq) in enumerate(rows,1):
        rate_formula=R.linked(key+'_rate') if duplicate else (R.product('database_usd_month','fx_php_usd') if key=='database' else None)
        annual=R.linked(key+'_annual') if duplicate else (lambda r,k=key,f=freq:f'=IF(ISNUMBER({r(k+"_rate")}),{r(k+"_rate")}*{f},"")')
        R.table_row(scene,x,y+i*19,[117,58,53],19,[name,'',''],ids=[None,prefix+key+'_rate',prefix+key+'_annual'],
          inputs=[False,not duplicate and key!='database',False],formulas=[None,rate_formula,annual],
          aligns=['left','right','right'],money=[False,True,True],size=8.6)
    R.table_row(scene,x,y+7*19,[175,53],20,['TOTAL',''],True,ids=[None,prefix+'recurring_total'],
      formulas=[None,R.linked('recurring_total') if duplicate else R.line_sum([key+'_annual' for _,key,_ in rows])],aligns=['left','right'],money=[False,True])

def basis_scene():
    s=R.Scene('Input Details','Planning Basis and Sources')
    s.label(35,38,525,27,'PLANNING BASIS — EDITABLE ESTIMATES (PHP)',13,True)
    s.label(35,76,525,43,'These financial values and task timings are illustrative estimates, not actual company expenses. Only the green inventory quantities are company data. A three-month build is assumed; project fees cover part-time services.',10.5,pad=0)
    R.table_row(s,35,134,[265,80,180],23,['Supporting input','Estimate','Unit / basis'],True,R.GRAY)
    rows=[('Training duration','training_days','days'),('Cloud deployment setup quantity','server_qty','setup service'),
      ('External backup drive quantity','drive_qty','one drive'),('Open-source / SSL package quantity','license_qty','free package'),
      ('Planning exchange rate','fx_php_usd','PHP per USD'),('Hosted backend subscription','database_usd_month','USD per month [S1]'),
      ('Administrative time valuation','hourly_labor_rate','PHP per person-hour'),
      ('Existing inventory-monitoring time','existing_labor1_hours','person-hours per month'),
      ('Existing report-preparation time','existing_labor2_hours','person-hours per month'),
      ('Existing reconciliation / follow-up time','existing_labor3_hours','person-hours per month'),
      ('Proposed inventory-monitoring time','proposed_labor1_hours','person-hours per month'),
      ('Proposed report-preparation time','proposed_labor2_hours','person-hours per month'),
      ('Proposed reconciliation / follow-up time','proposed_labor3_hours','person-hours per month')]
    for i,(label,key,unit) in enumerate(rows):
        R.table_row(s,35,157+i*20,[265,80,180],20,[label,'',unit],ids=[None,key,None],inputs=[False,True,False],money=[False,key=='hourly_labor_rate',False],aligns=['left','right','left'])
    s.label(35,433,525,54,'Cost boundaries: coding/testing is covered by the development retainers; PHP 1,000 setup covers account provisioning and deployment checks. The backup drive is a separate offline export copy. Hosted database storage and bundled database backups are paid once; frontend hosting, open-source licenses and SSL add no separate fee.',9.8,pad=0)
    s.label(35,496,525,56,'Labor basis: estimated monthly task time falls from 80 to 38 person-hours. The 42 hours released each month equal 504 hours/year; at PHP 150/hour, productivity value is PHP 75,600/year. This does not reduce payroll. Supplies savings are PHP 4,690/year; after PHP 31,800 recurring costs, net cash impact is negative PHP 27,110/year.',9.8,pad=0)
    s.label(35,563,525,18,'BENCHMARKS AND SOURCES — CHECKED 24 OCTOBER 2026',10,True)
    lines=[
      '[S1] Supabase Pro: USD 25/month starting price, within stated quotas. https://supabase.com/pricing',
      '[S2] Cloudflare Pages Free: static hosting limits apply. https://developers.cloudflare.com/pages/platform/limits/',
      '[S3] Registrar: no-markup domain pricing and included SSL. PHP 900 is an allowance. https://www.cloudflare.com/domains/',
      '[S4] Philippine wage context only; company wages are unknown. https://nwpc.dole.gov.ph/ncr/',
      '[S5] Official FX reference. PHP 65/USD is a rounded planning conversion. https://www.bsp.gov.ph/SitePages/Statistics/exchangerate.aspx',
      'Other rates, supplies quantities, hardware allowance and time reductions are analyst estimates, not quotations. Training includes a basic two-day orientation; any additional paid participant overtime needs a revised budget. Confirm taxes, actual usage and remaining development scope before spending.'
    ]
    for i,text in enumerate(lines):s.label(35,590+i*30,525,28,text,8.8,pad=0)
    s.label(35,801,480,14,'PLANNING ESTIMATES; COMPANY SOURCE: GREEN INVENTORY TABLE ONLY',8.5)
    s.label(521,801,44,14,'5',10,False,'right')
    return s

def configure():
    R.DOCUMENT_SCENE_COUNT=5
    R.DOCUMENT_TITLE='SiteTrace Cost-Benefit Worksheets — Realistic Planning Estimates'
    R.DOCUMENT_SUBJECT='Estimated financial scenario; green-table inventory quantities remain the only company data'
    R.WORKBOOK_DESCRIPTION='User-authorized illustrative planning estimates. Company quantities only from the green inventory image; no system demo data used.'
    R.INPUT_COMMENT='Editable planning estimate, not an actual company bill or wage. Replace with a verified quote or measured value when available. Zero means an explicit free/bundled/not-required item in this scenario.'
    R.PRESENTATION_NOTES='Financial figures are illustrative, editable planning estimates authorized by the user. They are not company actuals. Only the green inventory image supplies company quantities. Labor savings represent released capacity, not payroll reduction. Sources and timing assumptions appear on the final page.'
    scenes=R.SCENES
    for s in scenes[:3]:
        alter_text(s,'UNIT 2: Systems Planning','UNIT 2: Systems Planning | PLANNING ESTIMATES (PHP)')
    first,second,benefits,inventory=scenes[:4]
    alter_text(first,'Assuming that the project development duration is ______ MONTHS','Assumed project development duration: 3 MONTHS (part-time project fees)')
    alter_text(first,'Salary/Month','Fee/Month')
    for old,new in [('System Analyst','System Analyst (PT)'),('Programmer','Programmer (PT)'),('Encoder','Encoder (PT)'),('Researcher','Research / QA (PT)')]:alter_text(first,old,new)
    alter_text(first,'User Training','User Training (2 days)')
    for s in [first,second]:
        alter_text(s,'Database Server','Cloud Database Setup')
        alter_text(s,'Licensed Software /\nVendor Platform','Open-source Tools /\nFree SSL')
        cutoff=410 if s is first else 80
        s.blocks=[b for b in s.blocks if not(b.id and any(b.id.endswith(x) for x in ['maintenance_rate','maintenance_annual','communication_rate','communication_annual','renewal_rate','renewal_annual','recurring_total']))
          and not(cutoff<=b.y<=cutoff+76 and b.x>=R.RIGHT and (b.text in ['Description','Unit Cost','Annually','Software Maintenance','Incremental Communication','Software License Renewal','TOTAL']))]
        recurring(s,R.RIGHT,cutoff,s is second)
        # Replace the decorative empty square with a useful estimate/currency marker.
        marker=next(o for o in s.outlines if o.w==53)
        s.label(marker.x+1,marker.y+4,marker.w-2,marker.h-8,'PHP\nEST.',10,True,'center',pad=0)
    old='Financial amounts and development duration are blank because the green company inventory table does not supply them. Monthly maintenance/communication and annual renewal are worksheet billing units, not purchased services.'
    alter_text(first,old,'All financial figures are planning estimates in PHP, not actual company expenses. Maintenance, hosting and communication rates are monthly; renewals are annual. Free or bundled items are explicitly shown as 0.00. See Planning Basis for assumptions and sources.')
    alter_text(second,'This table follows the Breakdown of One Time Cost Worksheet in the supplied academic reference. The total one-time cost of PHP __________ will serve as the initial investment, and the annual recurring cost of PHP __________ will represent additional system expenses. These amounts are left blank because company financial data has not been supplied.',
      'The estimated one-time investment is PHP 114,500.00: development PHP 107,000.00, user training PHP 4,000.00, and deployment/backup hardware PHP 3,500.00. Estimated annual recurring cost is PHP 31,800.00. These figures form a planning scenario and are not actual company spending or approved vendor quotations.')
    second.label(44,606,507,76,'Estimated annual operational value is PHP 80,290.00, comprising PHP 4,690.00 in supplies cash savings and PHP 75,600.00 in released staff capacity. The capacity value is not payroll savings. Supplies savings alone do not cover the PHP 31,800.00 yearly running cost; the net cash requirement remains PHP 27,110.00 per year.',10.5,pad=0)
    for b in benefits.blocks:
        if b.text=='Personnel':b.text='Activity'
        if b.text in ['Folders','Ballpen','Bond Papers','Carbon Papers','Record Book','Vouchers','Ink','USB']:
            b.text={'Folders':'Folders (bundle)','Ballpen':'Ballpen (box)','Bond Papers':'Bond Papers (ream)',
              'Carbon Papers':'Carbon Paper (pack)','Record Book':'Record Book (pc)',
              'Vouchers':'Vouchers (pad)','Ink':'Ink (bottle)','USB':'USB (piece)'}[b.text]
        if b.text=='Salary/\nMonth':b.text='Allocated Cost/\nMonth'
        if b.text in ['Inventory personnel','Reporting personnel','Accountability personnel']:
            b.text={'Inventory personnel':'Inventory monitoring','Reporting personnel':'Report preparation','Accountability personnel':'Reconciliation / follow-up'}[b.text]
        if b.id and b.id.endswith('_monthly') and '_labor' in b.id:
            b.input=False;b.formula=R.product(b.id.replace('_monthly','_hours'),'hourly_labor_rate')
        if isinstance(b.text,str):
            b.text=b.text.replace('B. LABOR COST','B. LABOR AND PRODUCTIVITY VALUE')
            b.text=b.text.replace('Labor Cost Reduction = Labor Cost of the Existing System − Labor Cost of the Proposed System','Productivity Value = Existing Allocated Labor Cost − Proposed Allocated Labor Cost')
            b.text=b.text.replace('Labor Cost Reduction','Labor Productivity Value')
            b.text=b.text.replace('Total Cost Reduction = Labor Productivity Value + Supplies Cost Reduction','Total Annual Tangible Value = Productivity Value + Supplies Cash Savings')
            if b.text=='Total Cost Reduction':b.text='Total Annual Tangible Value'
    alter_text(benefits,'This table follows the Breakdown of Tangible Benefit Worksheet in the supplied reference. Annual supplies cost reduction and labor cost reduction remain blank because the green company inventory table provides no supplies spending, salaries or staffing figures. A reduction in employee headcount or payroll has not been assumed.',
      'Estimated annual supplies spending falls from PHP 7,500.00 to PHP 2,810.00, a potential cash saving of PHP 4,690.00. Allocated labor falls from PHP 144,000.00 to PHP 68,400.00, releasing capacity valued at PHP 75,600.00 per year. Combined tangible value is PHP 80,290.00; no salary or headcount reduction is assumed.')
    alter_text(benefits,'The 463 units shown as missing or without reports identify an equipment reconciliation concern. Their monetary value and any equipment-related savings remain blank because replacement prices and confirmed loss or recovery records were not supplied.',
      'No equipment-loss or recovery saving is credited (PHP 0.00). The 463 missing or unreported units remain a reconciliation concern, not confirmed losses. A USB purchase is unnecessary in this scenario because cloud reporting and the separately budgeted backup drive are available; its subtotal is PHP 0.00.')
    # Preserve unknown company facts, but display a dash so every table cell is meaningful.
    inventory.blocks=[b for b in inventory.blocks if not(602<=b.y<620)]
    for b in inventory.blocks:
        if b.id and (b.id.startswith('inventory_need') or b.id.startswith('inventory_gap')) and b.text=='':b.text='—'
    alter_text(inventory,'Blank cells are retained as blanks. Only three categories have supplied requirements and reported gaps. The 463 units are not automatically confirmed physical losses. “CATWALK FLATFORM SCAFFOLDING” preserves the spelling in the green source image.',
      'A dash means not reported in the source, not zero. Only three categories have supplied requirements and gaps. The 463 units are not confirmed physical losses. “CATWALK FLATFORM SCAFFOLDING” preserves the source spelling; unused row 27 is omitted.')
    R.SCENES=scenes[:4]+[basis_scene()]
    used=set()
    for s in R.SCENES:
        for b in s.blocks:
            if b.input:
                assert b.id in INPUTS,(s.name,b.id)
                b.text=INPUTS[b.id];used.add(b.id)
    assert used==set(INPUTS),('unused planning inputs',set(INPUTS)-used)

def fill_formula_values(manifest,values):
    id_values={}
    for key,(sheet,address) in manifest['ids'].items():
        coord=address.replace('$','');v=values.get((sheet.upper(),coord))
        if isinstance(v,(int,float)):id_values[key]=v
    for scene in R.SCENES:
        for block in scene.blocks:
            if block.formula:
                assert block.id in id_values,('unresolved formula',scene.name,block.id)
                block.text=id_values[block.id]
    R.EXCEL_CACHE=id_values
    return id_values

def validate(path,manifest,figures,pdf,docx,pptx):
    from openpyxl import load_workbook
    from lxml import etree
    from pptx import Presentation
    data=load_workbook(path,data_only=True);formulas=load_workbook(path,data_only=False)
    def get(key):
        sheet,cell=manifest['ids'][key];return data[sheet][cell.replace('$','')].value
    computed={k:get(k) for k in EXPECTED if k not in ['annual_net_cash','annual_net_economic']}
    checks={
      'all_financial_inputs_have_estimates':all(isinstance(formulas[v['sheet']][v['cell']].value,(int,float)) for v in manifest['inputs']),
      'all_financial_display_cells_are_numeric':all(isinstance(data[v['sheet']][v['cell']].value,(int,float)) for v in manifest['financial_cells']),
      'all_formula_caches_are_numeric':all(isinstance(data[v['sheet']][v['cell']].value,(int,float)) for v in manifest['formulas']),
      'all_totals_match_independent_calculation':all(math.isclose(v,EXPECTED[k],abs_tol=1e-9) for k,v in computed.items()),
      'no_cloud_service_double_count':get('database_annual')==19500 and get('frontend_annual')==0 and get('renewal_annual')==0 and get('recurring_total')==31800,
      'no_unnecessary_USB_purchase':get('proposed_supply5_qty')==0 and get('proposed_supply5_subtotal')==0,
      'payroll_not_reduced':get('labor_reduction')==75600 and EXPECTED['annual_net_cash']==-27110,
      'green_inventory_quantity_unchanged':sum(get('inventory_qty'+str(i)) for i in range(1,27))==758,
      'green_inventory_gaps_unchanged':sum(get('inventory_gap'+str(i)) for i in [2,17,26])==463,
      'source_unknowns_display_as_not_reported':sum(get('inventory_need'+str(i))=='—' for i in range(1,27))==23,
      'A4_portrait_print_layout':all(str(s.page_setup.paperSize)=='9' and s.page_setup.orientation=='portrait' and s.sheet_properties.pageSetUpPr.fitToPage for s in formulas),
      'no_external_workbook_links':not formulas._external_links,
      'page_geometry_clear':not R.validate_geometry(),
    }
    for i,(name,qty,need,gap) in enumerate(R.INVENTORY,1):
        assert get('inventory_qty'+str(i))==qty
        for label,value in [('need',need),('gap',gap)]:assert get(f'inventory_{label}{i}')==('—' if value is None else value)
    # A future missing estimate must not silently turn the investment into zero.
    with tempfile.TemporaryDirectory(prefix='sitettrace-estimate-check-') as folder:
        missing=load_workbook(path)
        sn,addr=manifest['ids']['dev1_rate'];missing[sn][addr.replace('$','')]=None
        target=Path(folder)/'missing.xlsx';missing.save(target)
        answer=R.evaluate_formulas(target)
        sn,addr=manifest['ids']['one_time_total']
        checks['missing_input_guard_preserved']=answer[(sn.upper(),addr.replace('$',''))]==''
    text={}
    for path_ in [path,docx,pptx]:
        with zipfile.ZipFile(path_) as z:
            assert z.testzip() is None
            for file in z.namelist():
                if file.endswith('.xml'):etree.fromstring(z.read(file))
            if path_==docx:
                root=etree.fromstring(z.read('word/document.xml'))
                ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','v':'urn:schemas-microsoft-com:vml'}
                text['Word']=' '.join(root.xpath('//w:t/text()',namespaces=ns))
                checks['Word_five_editable_pages']=len(root.xpath('//w:br[@w:type="page"]',namespaces=ns))==4 and len(root.xpath('//v:textbox',namespaces=ns))>100
    prs=Presentation(pptx);checks['PowerPoint_five_editable_slides']=len(prs.slides)==5 and not any(x.shape_type==13 for s in prs.slides for x in s.shapes)
    text['PowerPoint']=' '.join(sh.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    document=R.pymupdf.open(pdf);text['PDF']=' '.join(p.get_text() for p in document)
    checks['PDF_five_A4_pages']=len(document)==5 and all(abs(p.rect.width-R.W)<1 and abs(p.rect.height-R.H)<1 for p in document)
    checks['matching_financial_figures_in_all_formats']=all(all(number in t for number in ['114,500.00','31,800.00','80,290.00','75,600.00','4,690.00']) for t in text.values())
    checks['estimates_identified_in_all_editions']=all('PLANNING ESTIMATES' in t for t in text.values())
    obsolete=['financial amounts are blank','remain blank because','PHP __________','financial data has not been supplied','3 MONTHS (full-time)','BuildRight','Casa Buena','Metropolis','San Gabriel','178,500.00','20,500.00']
    checks['no_blank_version_prose_or_demo_facts']=not any(s in t for s in obsolete for t in text.values())
    overflow=[]
    for i,p in enumerate(document):
        for b in p.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                for span in line['spans']:
                    x0,y0,x1,y1=span['bbox']
                    if x0<0 or y0<0 or x1>R.W+.5 or y1>R.H+.5:overflow.append([i+1,span['text']])
    checks['PDF_text_inside_page']=not overflow
    for s in R.SCENES:
        for b in s.blocks:
            for line in R.block_lines(b):
                width=R.pdfmetrics.stringWidth(line,'B' if b.bold else 'N',b.size)
                assert width<=b.w-2*b.pad+.8,(s.name,line,width,b.w-2*b.pad)
    assert all(checks.values()),{'checks':checks,'overflow':overflow}
    return {'checks':checks,'prepared_date':'2026-10-24','timezone':'Asia/Manila',
      'scenario':'Illustrative analyst estimates authorized by user; no company financial actuals or system demo records used.',
      'company_source':'Green inventory only: 26 categories, 758 recorded units, 463 missing or unreported units.',
      'financial_totals':EXPECTED,'formula_count':len(manifest['formulas']),'editable_estimate_count':len(manifest['inputs']),
      'labor_basis':{'hourly_rate':150,'monthly_existing_hours':80,'monthly_proposed_hours':38,'annual_hours_released':504,'annual_productivity_value':75600,'payroll_cash_reduction_assumed':0},
      'sources':SOURCES,'PDF_pages':5,'PowerPoint_slides':5,
      'rendering_limit':'Native Office automation was unavailable in this session (COM 80070520). OOXML packages, formula calculations/caches, common vector geometry and PDF rendering checked. Minor native Office layout differences may occur.'}

def preview(pdf):
    from PIL import Image,ImageDraw,ImageFont
    names=['01-one-time-cost','02-cost-explanation','03-tangible-benefits','04-green-company-inventory','05-planning-basis']
    p=R.pymupdf.open(pdf);contact=Image.new('RGB',(1575,485),'#E8EBEF')
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',12)
    for i,page in enumerate(p):
        file=VIEWS/(names[i]+'.png');page.get_pixmap(matrix=R.pymupdf.Matrix(2,2)).save(file)
        im=Image.open(file).convert('RGB');im.thumbnail((300,425));contact.paste(im,(i*315+7,10))
        ImageDraw.Draw(contact).text((i*315+8,450),names[i],font=font,fill='black')
    contact.save(VAL/'reference-recreation-contact.png')

def documentation(report):
    text='''# SiteTrace cost-benefit worksheets — populated planning estimates

The current files fill the financial tables with user-authorized, realistic **planning estimates**. They are not actual company salaries, expenses, quotations or measured benefits. Company data still comes only from the green inventory image: **26 categories, 758 recorded units and 463 missing or unreported units**. No system demonstration records are used.

## Current files

- `SiteTrace_Cost_Benefit_Report.pdf`: five A4 pages, with the three reference-style worksheet views, the green company inventory, and the planning basis/sources.
- `SiteTrace_Cost_Benefit_Report.docx`: the same five pages as editable Word text and vector shapes.
- `SiteTrace_Cost_Benefit_Worksheets.pptx`: five editable portrait PowerPoint slides.
- `SiteTrace_Cost_Benefit_Analysis.xlsx`: five editable, printable worksheets with cached calculated figures and automatic recalculation.
- `recreated-views/`: high-resolution images of all five pages.

## Estimated figures

| Item | PHP |
|---|---:|
| Part-time development services | 107,000.00 |
| Two-day training and materials | 4,000.00 |
| Deployment setup and backup drive | 3,500.00 |
| **One-time investment** | **114,500.00** |
| **Annual recurring cost** | **31,800.00** |
| Existing annual supplies | 7,500.00 |
| Proposed annual supplies | 2,810.00 |
| **Potential annual supplies cash saving** | **4,690.00** |
| Released staff capacity value | 75,600.00 |
| **Annual tangible operational value** | **80,290.00** |
| Operational value less recurring cost (includes noncash capacity) | 48,490.00 |
| **Cash impact from supplies saving less recurring cost** | **−27,110.00** |

The labor comparison values time allocated to inventory work; it does **not** assume salary cuts, layoffs or lower headcount. Estimated task time falls from 80 to 38 person-hours per month at PHP 150/hour, releasing 504 hours/year. Staff can use that capacity for other work. Only supplies spending is counted as a potential cash saving, so the system still needs an estimated PHP 27,110/year cash operating allocation before any separate proven payroll or equipment saving. The initial investment is additional.

## Why the estimates are plausible

- Development is a three-month, part-time project allowance: analyst PHP 10,000/month for 3 months; programmer PHP 20,000/month for 3 months; encoding PHP 5,000 for 1 month; research/QA PHP 6,000/month for 2 months. These are capped part-time service retainers, not full-time employee salaries. Full small-system build/implementation scope is costed; exclude already-incurred development before using the workbook for a remaining-work purchase decision.
- Training is two trainer days at PHP 1,500/day plus PHP 1,000 materials. The scenario assumes attendance during normal paid hours; separately paid overtime would require an added allowance.
- One-time production setup is PHP 1,000, independent offline backup drive allowance is PHP 2,500, and open-source/SSL license fees are zero. No unnecessary physical database server is purchased.
- Hosted backend is budgeted at Supabase Pro's published USD 25/month starting level, converted at a rounded planning PHP 65/USD: PHP 1,625/month. Usage is assumed to stay inside the included plan limits. Database storage and bundled daily database backups are not charged again. This is not a claim about an existing company subscription.
- Static frontend hosting uses Cloudflare Pages Free within its published limits. Maintenance is PHP 750/month; incremental communication PHP 200/month; domain allowance PHP 900/year. Free frontend hosting, SSL and open-source licenses explicitly show 0.00. Domain and hardware allowances are estimates, not vendor quotations. Confirm actual bank conversion, taxes, usage and any overage charges before buying.
- Annual paper, stationery and ink quantities are modest planning allowances. USB quantity is zero because cloud reports and the separately budgeted backup drive cover this scenario. Prices and task-time reductions are assumptions, not findings from the green inventory.
- NWPC wage information is a general Philippine costing benchmark only. It is not company payroll evidence or a legal wage-compliance assessment. No equipment recovery, theft prevention, reduced loss or avoided procurement is monetized. The 463-unit reporting gap is not a confirmed financial loss.

## Sources checked 24 October 2026

- [Supabase pricing](https://supabase.com/pricing): Pro starts at USD 25/month, with published quotas and database-backup retention.
- [Cloudflare Pages limits](https://developers.cloudflare.com/pages/platform/limits/): Free static hosting subject to limits.
- [Cloudflare Registrar](https://www.cloudflare.com/domains/): domain pricing model and included DNS/SSL; the PHP 900 domain allowance is an estimate.
- [NWPC NCR](https://nwpc.dole.gov.ph/ncr/): general local wage context; project retainers and PHP 150/hour valuation are analyst estimates.
- [BSP exchange-rate statistics](https://www.bsp.gov.ph/SitePages/Statistics/exchangerate.aspx): official reference resource; PHP 65/USD is the model's rounded planning rate.

## Editing and verification

All financial input cells now have numbers; genuine free/bundled items are 0.00. Supply descriptions specify bundles, boxes, reams, packs, pieces, pads or bottles so each unit price is clear. Replace these estimates with actual company rates when available. The Input Details sheet controls training duration, hardware quantities, hosting FX and six monthly activity-hour assumptions. Labor allocations and hosted-backend PHP cost calculate from those inputs. Duplicate cost-view cells link to the first worksheet, rather than creating another charge. Excel formula guards continue to display blank if an essential estimate is deleted. Narrative paragraphs describe the saved baseline and should be refreshed when a scenario is changed; calculated value cells recalculate automatically.

The green inventory remains separate. A dash means not reported, not zero. No quantities or shortages were invented for the other categories. The unused source row 27 is omitted. The original source spelling “CATWALK FLATFORM SCAFFOLDING” is retained.

`validation/validation-report.json` records independent total checks, complete financial cells, formula caches, cloud-service cost boundaries, green-data preservation, missing-input protection, consistent figures in all formats, XML package integrity and A4 layout checks. The Word/PowerPoint editions use editable text/vector shapes, not pasted screenshots. Native Office automation was unavailable (COM 80070520); PDF and common scene geometry were checked. Minor native Office differences may occur.

`archive/blank-reference-edition/` preserves the previous blank version. `archive/initial-analysis/` is an older, superseded analysis. Open the current files in this directory to see the completed estimates. Application source, database schemas and production settings remain untouched.

Rebuild this populated scenario with Python 3.11 and `source/requirements.txt`: `python source/build_estimated_cba.py`. It uses `build_reference_recreation.py` for the reference-style layout and `build_estimated_cba.py` for the explicit input estimates. Rebuilding restores this planning scenario, so back up later workbook edits. Word/PowerPoint/PDF are static snapshots; refresh those editions alongside any revised Excel figures.
'''
    (OUT/'README.md').write_text(text,encoding='utf-8')
    (VAL/'planning-estimates.json').write_text(json.dumps({'classification':'Illustrative planning estimates, not company actuals','inputs':INPUTS,'independent_totals':EXPECTED,'sources':SOURCES},indent=2),encoding='utf-8')
    provenance=json.loads((VAL/'company-data-source.json').read_text(encoding='utf-8'))
    provenance['policy']='Green image supplies company inventory only. User-authorized financial estimates are analyst projections; reference photos supply layout; no system demo records are used.'
    (VAL/'company-data-source.json').write_text(json.dumps(provenance,indent=2),encoding='utf-8')
    report['artifacts']=[{'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [OUT/'SiteTrace_Cost_Benefit_Analysis.xlsx',OUT/'SiteTrace_Cost_Benefit_Report.docx',OUT/'SiteTrace_Cost_Benefit_Report.pdf',OUT/'SiteTrace_Cost_Benefit_Worksheets.pptx']]
    (VAL/'validation-report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')

def main():
    configure();R.fit_padding();assert not R.validate_geometry(),R.validate_geometry()
    # First evaluate Excel's actual cross-sheet formulas, then cache those results
    # and reuse them in every printed/editable document edition.
    workbook,manifest=R.excel_workbook();print('All planning inputs populated; calculating formulas',flush=True)
    values=R.evaluate_formulas(workbook);figures=fill_formula_values(manifest,values)
    workbook,manifest=R.excel_workbook();R.fit_padding()
    pdf=OUT/'SiteTrace_Cost_Benefit_Report.pdf';R.render_pdf(pdf,R.SCENES)
    pptx=R.powerpoint();docx=R.word_document();print('Five-page document editions generated',flush=True)
    report=validate(workbook,manifest,figures,pdf,docx,pptx)
    preview(pdf);documentation(report)
    print(json.dumps({'financial_estimates':EXPECTED,'formulas':report['formula_count'],'pages':5,'all_checks_passed':all(report['checks'].values())},indent=2),flush=True)

if __name__=='__main__':main()

