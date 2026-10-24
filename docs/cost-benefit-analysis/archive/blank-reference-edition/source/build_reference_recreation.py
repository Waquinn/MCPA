"""Recreate the three supplied worksheet photographs using only green-table data.

No application module, demo record, repository documentation, financial estimate,
reference-photo example amount or assumed development duration is read/imported.
The same vector layout generates PDF, editable PowerPoint and editable Word shapes.
Excel contains corresponding editable cells and guarded formulas.
"""
from pathlib import Path
from dataclasses import dataclass, field
from datetime import datetime, timezone
from xml.sax.saxutils import escape
import csv, hashlib, json, math, re, tempfile, zipfile
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
import xlsxwriter
from openpyxl import load_workbook
import pymupdf

OUT=Path(__file__).resolve().parents[1]
VAL=OUT/'validation'
VIEWS=OUT/'recreated-views'
W,H=A4
MAROON='6A1839'; BLUE='D9E6F3'; GRAY='EFEFEF'; GREEN='B5D6A4'; LIGHT_GREEN='DCEAD2'
FONT_DIR=Path('C:/Windows/Fonts')
NARROW=(FONT_DIR/'arialn.ttf').exists()
for key,name in [('N','arialn.ttf' if NARROW else 'arial.ttf'),('B','arialnb.ttf' if NARROW else 'arialbd.ttf'),('I','arialni.ttf' if NARROW else 'ariali.ttf')]:
    pdfmetrics.registerFont(TTFont(key,str(FONT_DIR/name)))
FONT_FACE='Arial Narrow' if NARROW else 'Arial'

INVENTORY=[
 ('SHOVEL (PALA, LAPAD)',19,None,None),('SHOVEL (PALA, BILOG)',23,42,19),
 ('BARETA',13,None,None),('ITAK',4,None,None),('PIKO',4,None,None),
 ('PLASTIC DRUM 200L',14,None,None),('IBC WATER TANK 1000L',5,None,None),
 ('WHEEL BARROW',4,None,None),('OXYGEN TANK',3,None,None),('GAS TANK',5,None,None),
 ('ELECTRIC KETTLE',2,None,None),('EXTENSION',10,None,None),
 ('DOUBLE BURNER STOVE',8,None,None),('LABABO',5,None,None),('MASO',4,None,None),
 ('GARDEN HOSE',5,None,None),('GI PIPES',225,300,75),('BASE JACK',27,None,None),
 ('U-JACK',56,None,None),('BUCKETS',39,None,None),
 ('CATWALK FLATFORM SCAFFOLDING',33,None,None),('H-FRAME SCAFFOLDING',64,None,None),
 ('BRACES SCAFFOLDING',41,None,None),('WATER-PUMP',2,None,None),
 ('GARDEN SPRINKLER',2,None,None),('SWIVEL CLAMP',141,510,369)]
assert len(INVENTORY)==26
assert sum(x[1] for x in INVENTORY)==758
assert sum(x[3] or 0 for x in INVENTORY)==463

@dataclass
class Block:
    x:float;y:float;w:float;h:float;text:object='';size:float=9
    bold:bool=False;align:str='left';fill:str|None=None;border:float=0
    color:str='000000';pad:float=2;id:str|None=None;input:bool=False
    money:bool=False;formula:object=None;note:str=''

@dataclass
class Outline:
    x:float;y:float;w:float;h:float;width:float=1.2

@dataclass
class Scene:
    name:str;title:str;blocks:list=field(default_factory=list);outlines:list=field(default_factory=list)
    def label(self,x,y,w,h,text,size=10,bold=False,align='left',**kw):
        b=Block(x,y,w,h,text,size,bold,align,**kw);self.blocks.append(b);return b
    def box(self,x,y,w,h,text='',size=9,bold=False,align='left',**kw):
        b=Block(x,y,w,h,text,size,bold,align,border=.5,**kw);self.blocks.append(b);return b

def line_sum(ids):
    def f(r):
        cells=','.join(r(i) for i in ids)
        return f'=IF(COUNT({cells})={len(ids)},SUM({cells}),"")'
    return f
def product(a,b):return lambda r:f'=IF(COUNT({r(a)},{r(b)})=2,{r(a)}*{r(b)},"")'
def delta(a,b):return lambda r:f'=IF(COUNT({r(a)},{r(b)})=2,{r(a)}-{r(b)},"")'
def linked(a):return lambda r:f'=IF(ISNUMBER({r(a)}),{r(a)},"")'

def table_row(s,x,y,widths,h,texts,bold=False,fill=None,aligns=None,ids=None,inputs=None,formulas=None,money=None,size=9):
    for i,(width,text) in enumerate(zip(widths,texts)):
        s.box(x,y,width,h,text,size,bold,(aligns or ['left']*len(widths))[i],
          fill=fill,id=(ids or [None]*len(widths))[i],input=(inputs or [False]*len(widths))[i],
          formula=(formulas or [None]*len(widths))[i],money=(money or [False]*len(widths))[i])
        x+=width

def header(s,page,benefit=False):
    s.label(44,44,507,18,'UNIT 2: Systems Planning',10,True,fill=MAROON,color='FFFFFF',pad=4)
    if benefit:s.label(51,70,480,20,'•  Benefit',14,True)
    s.label(44,801,455,14,'IT303: SYSTEMS ANALYSIS AND DESIGN',8.5)
    s.label(507,801,44,14,str(page),10,False,'right')

LEFT=54;RIGHT=311;LW=224;RW=228

def development_table(s):
    s.label(LEFT,150,LW,20,'A. Development Cost',11)
    widths=[104,61,24,35]
    table_row(s,LEFT,178,widths,19,['Description','Salary/Month','Mo.','Sub_Total'],True,GRAY,size=8.7)
    for i,name in enumerate(['System Analyst','Programmer','Encoder','Researcher'],1):
        table_row(s,LEFT,197+(i-1)*19,widths,19,[name,'','',''],
          ids=[None,f'dev{i}_rate',f'dev{i}_months',f'dev{i}_subtotal'],inputs=[False,True,True,False],
          aligns=['left','right','center','right'],formulas=[None,None,None,product(f'dev{i}_rate',f'dev{i}_months')],money=[False,True,False,True],size=9)
    table_row(s,LEFT,273,[189,35],20,['TOTAL',''],True,
      ids=[None,'development_total'],formulas=[None,line_sum([f'dev{i}_subtotal' for i in range(1,5)])],aligns=['left','right'],money=[False,True],size=9)

def training_table(s):
    s.label(LEFT,312,LW,20,'B. User Training',11)
    widths=[104,70,50]
    table_row(s,LEFT,341,widths,19,['Description','Honoraria/day','Sub_Total'],True,GRAY,size=8.9)
    table_row(s,LEFT,360,widths,19,['User Training','',''],ids=[None,'training_rate','training_subtotal'],
      inputs=[False,True,False],formulas=[None,None,product('training_rate','training_days')],aligns=['left','right','right'],money=[False,True,True])
    table_row(s,LEFT,379,widths,19,['Training Kit','',''],ids=[None,'training_kit','training_kit_subtotal'],
      inputs=[False,True,False],formulas=[None,None,linked('training_kit')],aligns=['left','right','right'],money=[False,True,True])
    table_row(s,LEFT,398,[174,50],20,['TOTAL',''],True,ids=[None,'training_total'],
      formulas=[None,line_sum(['training_subtotal','training_kit_subtotal'])],aligns=['left','right'],money=[False,True])

def hardware_table(s,x,y,duplicate=False):
    prefix='copy_' if duplicate else ''
    widths=[104,70,50]
    table_row(s,x,y,widths,19,['Description','Unit Cost','Sub_Total'],True,GRAY,size=9)
    row_defs=[('Database Server','server',19),('External Drive','drive',19),('Licensed Software /\nVendor Platform','license',36)]
    yy=y+19
    for label,key,height in row_defs:
        if duplicate:
            forms=[None,linked(f'{key}_rate'),linked(f'{key}_subtotal')];ins=[False]*3
        else:
            forms=[None,None,product(f'{key}_rate',f'{key}_qty')];ins=[False,True,False]
        table_row(s,x,yy,widths,height,[label,'',''],ids=[None,prefix+f'{key}_rate',prefix+f'{key}_subtotal'],
          inputs=ins,formulas=forms,aligns=['left','right','right'],money=[False,True,True]);yy+=height
    table_row(s,x,yy,[174,50],20,['TOTAL',''],True,ids=[None,prefix+'hardware_total'],
      formulas=[None,linked('hardware_total') if duplicate else line_sum(['server_subtotal','drive_subtotal','license_subtotal'])],aligns=['left','right'],money=[False,True])
    return yy+20

def formula_box(s,x,y,w,h):
    s.box(x,y,w,h,'Total One Time Cost = Total Development Cost + Total\nUser Training + Total Additional Hardware/Software',9.5,pad=5)

def recurring_table(s,x,y,duplicate=False):
    prefix='copy_' if duplicate else ''
    widths=[117,58,53]
    table_row(s,x,y,widths,19,['Description','Unit Cost','Annually'],True,GRAY,size=9)
    for i,(label,key,freq) in enumerate([('Software Maintenance','maintenance',12),('Incremental Communication','communication',12),('Software License Renewal','renewal',1)],1):
        forms=[None,linked(key+'_rate'),linked(key+'_annual')] if duplicate else [None,None,lambda r,k=key,f=freq:f'=IF(ISNUMBER({r(k+"_rate")}),{r(k+"_rate")}*{f},"")']
        table_row(s,x,y+i*19,widths,19,[label,'',''],ids=[None,prefix+key+'_rate',prefix+key+'_annual'],
          inputs=[False,not duplicate,False],formulas=forms,aligns=['left','right','right'],money=[False,True,True],size=8.8)
    table_row(s,x,y+76,[175,53],20,['TOTAL',''],True,ids=[None,prefix+'recurring_total'],
      formulas=[None,linked('recurring_total') if duplicate else line_sum(['maintenance_annual','communication_annual','renewal_annual'])],aligns=['left','right'],money=[False,True])

def one_time_scene():
    s=Scene('One Time Cost','Breakdown of One Time Cost Worksheet');header(s,1)
    s.outlines.append(Outline(44,82,507,580,1.5))
    s.label(54,91,485,19,'BREAKDOWN OF ONE TIME COST WORKSHEET',12,True,'center')
    s.label(54,117,485,20,'Assuming that the project development duration is ______ MONTHS',10.5)
    development_table(s);training_table(s)
    s.label(LEFT,440,LW,20,'C. ADDITIONAL HARDWARE/SOFTWARE',10.5)
    hardware_table(s,LEFT,470)
    formula_box(s,LEFT,610,LW,42)
    s.label(RIGHT,151,RW,29,'BREAKDOWN OF ONE TIME COST WORKSHEET',10.5,True,'center')
    rows=[('A. Development Cost','development_total',19),('B. User Training','training_total',19),('C. Additional\n    Hardware/Software','hardware_total',34),('TOTAL ONE TIME COST','one_time_total',21)]
    y=191
    for label,key,height in rows:
        calc=line_sum(['development_total','training_total','hardware_total']) if key=='one_time_total' else linked(key)
        table_row(s,RIGHT,y,[167,61],height,[label,''],key=='one_time_total',ids=[None,'summary_'+key if key!='one_time_total' else key],
          formulas=[None,calc],aligns=['left','right'],money=[False,True],size=9.3);y+=height
    s.box(RIGHT,y,RW,44,'Total One Time Cost = Total Development Cost + Total\nUser Training + Total Additional Hardware/Software',9.2,pad=4)
    s.label(RIGHT,389,RW,20,'RECURRING COST',11,True)
    recurring_table(s,RIGHT,410)
    s.outlines.append(Outline(485,598,53,36,.8))
    s.label(44,683,507,31,'Financial amounts and development duration are blank because the green company inventory table does not supply them. Monthly maintenance/communication and annual renewal are worksheet billing units, not purchased services.',9.5)
    return s

def explanation_scene():
    s=Scene('Cost Explanation','One-Time and Recurring Cost Explanation');header(s,2)
    # The second photograph repeats the lower part of the first cost figure.
    # Rebuild those rows with live links; do not create a second set of inputs.
    s.outlines.append(Outline(44,80,507,300,1.5))
    table_row(s,LEFT,80,[104,70,50],19,['Training Kit','',''],
      ids=[None,'copy_training_kit','copy_training_kit_subtotal'],formulas=[None,linked('training_kit'),linked('training_kit_subtotal')],aligns=['left','right','right'],money=[False,True,True])
    table_row(s,LEFT,99,[174,50],20,['TOTAL',''],True,ids=[None,'copy_training_total'],formulas=[None,linked('training_total')],aligns=['left','right'],money=[False,True])
    s.label(LEFT,151,LW,20,'C. ADDITIONAL HARDWARE/SOFTWARE',10.5)
    hardware_table(s,LEFT,179,True)
    formula_box(s,LEFT,318,LW,49)
    recurring_table(s,RIGHT,80,True)
    s.outlines.append(Outline(485,318,53,36,.8))
    s.label(44,405,507,77,'This table follows the Breakdown of One Time Cost Worksheet in the supplied academic reference. The total one-time cost of PHP __________ will serve as the initial investment, and the annual recurring cost of PHP __________ will represent additional system expenses. These amounts are left blank because company financial data has not been supplied.',11,pad=0)
    s.label(44,503,507,68,'The green company monitoring table records 758 units across 26 equipment categories. It identifies 19 round shovels, 75 GI pipes and 369 swivel clamps under “Missing or No Report,” totaling 463 units. These are reconciliation gaps; the table does not establish confirmed financial losses.',11,pad=0)
    return s

def supplies_table(s,x,y,proposed=False):
    prefix='proposed' if proposed else 'existing';widths=[79,49,49,61]
    s.box(x,y,238,13,'PROPOSED SYSTEM' if proposed else 'EXISTING SYSTEM',9.4,True,'center',fill=GRAY)
    table_row(s,x,y+13,widths,13,['Description','Qty','Price/Qty','Sub_Total'],True,GRAY,size=8.8)
    names=['Folders','Ballpen','Bond Papers','Record Book','USB','Ink'] if proposed else ['Folders','Ballpen','Bond Papers','Carbon Papers','Record Book','Vouchers','Ink']
    for i,label in enumerate(names):
        key=prefix+'_supply'+str(i+1)
        table_row(s,x,y+26+i*13,widths,13,[label,'','',''],ids=[None,key+'_qty',key+'_price',key+'_subtotal'],
          inputs=[False,True,True,False],formulas=[None,None,None,product(key+'_qty',key+'_price')],aligns=['left','center','right','right'],money=[False,False,True,True],size=8.8)
    bottom=y+26+7*13
    # The proposed table has six supplies rows; preserve the reference's offset.
    table_row(s,x,bottom,[177,61],15,['TOTAL',''],True,ids=[None,prefix+'_supplies_total'],
      formulas=[None,line_sum([prefix+'_supply'+str(i+1)+'_subtotal' for i in range(len(names))])],aligns=['left','right'],money=[False,True],size=8.8)

def labor_table(s,x,y,proposed=False):
    prefix='proposed' if proposed else 'existing';widths=[86,71,71]
    s.box(x,y,228,13,'PROPOSED SYSTEM' if proposed else 'EXISTING SYSTEM',9.3,True,'center',fill=GRAY)
    table_row(s,x,y+13,widths,25,['Personnel','Salary/\nMonth','Annual'],True,GRAY,size=9)
    for i,name in enumerate(['Inventory personnel','Reporting personnel','Accountability personnel'],1):
        key=prefix+'_labor'+str(i)
        table_row(s,x,y+38+(i-1)*15,widths,15,[name,'',''],ids=[None,key+'_monthly',key+'_annual'],inputs=[False,True,False],
          formulas=[None,None,lambda r,k=key:f'=IF(ISNUMBER({r(k+"_monthly")}),{r(k+"_monthly")}*12,"")'],aligns=['left','right','right'],money=[False,True,True],size=8.6)
    table_row(s,x,y+83,[157,71],15,['TOTAL',''],True,ids=[None,prefix+'_labor_total'],
      formulas=[None,line_sum([prefix+'_labor'+str(i)+'_annual' for i in range(1,4)])],aligns=['left','right'],money=[False,True])

def benefits_scene():
    s=Scene('Tangible Benefits','Breakdown of Tangible Benefit Worksheet');header(s,3,True)
    s.outlines.append(Outline(44,103,507,398,1.5))
    s.label(51,106,493,14,'Breakdown of Tangible Benefit Worksheet',10,True)
    s.label(51,120,493,15,'A. SUPPLIES COST',10,True)
    supplies_table(s,51,142);supplies_table(s,306,142,True)
    s.label(51,277,493,15,'Supplies Cost Reduction = Total Supplies of the Existing System − Total Supplies of the Proposed System',8.8,True)
    s.label(51,292,417,16,'Supplies Cost Reduction',9.3,True,fill=BLUE)
    s.label(468,292,76,16,'',9.3,True,'right',fill=BLUE,id='supplies_reduction',money=True,formula=delta('existing_supplies_total','proposed_supplies_total'))
    s.label(51,318,493,15,'B. LABOR COST',10,True)
    labor_table(s,51,337);labor_table(s,306,337,True)
    s.label(51,442,493,14,'Labor Cost Reduction = Labor Cost of the Existing System − Labor Cost of the Proposed System',8.8,True)
    s.label(51,456,417,15,'Labor Cost Reduction',9.3,True,fill=BLUE)
    s.label(468,456,76,15,'',9.3,True,'right',fill=BLUE,id='labor_reduction',money=True,formula=delta('existing_labor_total','proposed_labor_total'))
    s.label(51,475,493,12,'Total Cost Reduction = Labor Cost Reduction + Supplies Cost Reduction',8.8,True)
    s.label(51,487,417,12,'Total Cost Reduction',9.3,True,fill=BLUE)
    s.label(468,487,76,12,'',9.3,True,'right',fill=BLUE,id='total_reduction',money=True,formula=line_sum(['labor_reduction','supplies_reduction']))
    s.label(44,521,507,69,'This table follows the Breakdown of Tangible Benefit Worksheet in the supplied reference. Annual supplies cost reduction and labor cost reduction remain blank because the green company inventory table provides no supplies spending, salaries or staffing figures. A reduction in employee headcount or payroll has not been assumed.',11,pad=0)
    s.label(44,607,507,59,'The 463 units shown as missing or without reports identify an equipment reconciliation concern. Their monetary value and any equipment-related savings remain blank because replacement prices and confirmed loss or recovery records were not supplied.',10.5,pad=0)
    return s

def inventory_scene():
    s=Scene('Company Inventory','Hand Tools and Others Monitoring')
    s.label(30,53,535,29,'HAND TOOLS AND OTHERS MONITORING',15,True,fill='98C37F',pad=4)
    widths=[39,238,37,110,111];x=30;y=99
    table_row(s,x,y,widths,35,['ITEM\nNO.','POWER TOOLS','QTY.','NEED TO COMPLY','MISSING OR\nNO REPORT'],True,GREEN,aligns=['center']*5,size=9.5)
    for i,(name,qty,need,gap) in enumerate(INVENTORY,1):
        fill=GREEN if i%2==0 else 'FFFFFF'
        table_row(s,x,y+35+(i-1)*18,widths,18,[i,name,qty,need if need is not None else '',gap if gap is not None else ''],False,fill,
          ids=[None,None,f'inventory_qty{i}',f'inventory_need{i}',f'inventory_gap{i}'],aligns=['center']*5,size=9.1)
    table_row(s,x,y+35+26*18,widths,18,[27,'','','',''],False,'FFFFFF',aligns=['center']*5,size=9.1)
    s.label(30,646,535,18,'SOURCE CHECK: 26 populated categories; 758 recorded units; 463 missing or unreported units.',10,True)
    s.label(30,674,535,52,'Blank cells are retained as blanks. Only three categories have supplied requirements and reported gaps. The 463 units are not automatically confirmed physical losses. “CATWALK FLATFORM SCAFFOLDING” preserves the spelling in the green source image.',10,pad=0)
    s.label(30,743,535,31,'Source: user-supplied green “HAND TOOLS AND OTHERS MONITORING” image. Application records are demonstration data and are excluded.',9.5,pad=0)
    s.label(30,801,490,14,'COMPANY DATA SOURCE: GREEN INVENTORY TABLE ONLY',8.5)
    s.label(521,801,44,14,'4',10,False,'right')
    return s

def inputs_scene():
    s=Scene('Input Details','Supporting Inputs')
    s.label(35,45,525,26,'SUPPORTING INPUTS — LEAVE UNKNOWN VALUES BLANK',13,True)
    s.label(35,84,525,47,'These inputs complete quantities omitted from the reference layout. They are not supplied by the green company inventory table. No value or development duration is assumed.',10.5,pad=0)
    table_row(s,35,147,[270,80,175],24,['Input','Value','Unit / purpose'],True,GRAY)
    data=[('Training delivery duration','training_days','days'),('Database server quantity','server_qty','items / services'),('External drive quantity','drive_qty','items'),('Software / vendor platform quantity','license_qty','licenses / packages')]
    for i,(label,key,unit) in enumerate(data):
        table_row(s,35,171+i*25,[270,80,175],25,[label,'',unit],ids=[None,key,None],inputs=[False,True,False])
    s.label(35,301,525,58,'Financial inputs on all sheets are intentionally blank. For labor, enter the combined monthly cost for the personnel group shown, rather than assuming any employee count. Annualization uses 12 months as a unit conversion. Time saved does not automatically reduce payroll.',10.5,pad=0)
    return s

SCENES=[one_time_scene(),explanation_scene(),benefits_scene(),inventory_scene(),inputs_scene()]

def wrap_text(text,size,width,bold=False):
    font='B' if bold else 'N';lines=[]
    for paragraph in str(text).split('\n'):
        current=''
        for word in paragraph.split():
            candidate=(current+' '+word).strip()
            if current and pdfmetrics.stringWidth(candidate,font,size)>width:lines.append(current);current=word
            else:current=candidate
        lines.append(current)
    return lines

def block_lines(b):
    return wrap_text(b.text,b.size,max(1,b.w-2*b.pad),b.bold) if b.text!='' else []

def fit_padding():
    for s in SCENES:
        for b in s.blocks:
            lines=block_lines(b)
            if lines:b.pad=min(b.pad,max(.1,(b.h-len(lines)*b.size*1.12)/2-.3))

def validate_geometry():
    issues=[]
    for s in SCENES:
        for b in s.blocks:
            lines=block_lines(b)
            if len(lines)*b.size*1.12+2*b.pad>b.h+.6:issues.append([s.name,b.text,'text height',len(lines)*b.size*1.12+2*b.pad,b.h])
            if not(0<=b.x and 0<=b.y and b.x+b.w<=W+1 and b.y+b.h<=H):issues.append([s.name,b.text,'page bounds'])
    return issues

def render_pdf(path,scenes):
    c=canvas.Canvas(str(path),pagesize=A4,pageCompression=1)
    c.setTitle('SiteTrace Cost-Benefit Worksheets — Reference Recreation')
    c.setAuthor('SiteTrace Academic Documentation')
    for index,s in enumerate(scenes,1):
        c.bookmarkPage('page'+str(index));c.addOutlineEntry(s.title,'page'+str(index),0)
        for out in s.outlines:
            c.setStrokeColor(HexColor('#000000'));c.setLineWidth(out.width)
            c.rect(out.x,H-out.y-out.h,out.w,out.h,stroke=1,fill=0)
        for b in s.blocks:
            if b.fill or b.border:
                c.setLineWidth(b.border or .01)
                if b.fill:c.setFillColor(HexColor('#'+b.fill))
                c.setStrokeColor(HexColor('#000000'))
                c.rect(b.x,H-b.y-b.h,b.w,b.h,fill=bool(b.fill),stroke=bool(b.border))
            c.setFillColor(HexColor('#'+b.color));c.setFont('B' if b.bold else 'N',b.size)
            for i,line in enumerate(block_lines(b)):
                yy=H-b.y-b.pad-b.size*.91-i*b.size*1.12
                if b.align=='center':c.drawCentredString(b.x+b.w/2,yy,line)
                elif b.align=='right':c.drawRightString(b.x+b.w-b.pad,yy,line)
                else:c.drawString(b.x+b.pad,yy,line)
        c.showPage()
    c.save()

def excel_workbook():
    from xlsxwriter.utility import xl_rowcol_to_cell
    path=OUT/'SiteTrace_Cost_Benefit_Analysis.xlsx'
    book=xlsxwriter.Workbook(path)
    book.set_properties({'title':'SiteTrace Cost-Benefit Worksheets — Reference Recreation','author':'SiteTrace Academic Documentation','comments':'Only the green supplied inventory image supplies company data. Financial amounts intentionally blank.'})
    book.set_calc_mode('auto')
    grids={};ids={};formula_cells=[];input_cells=[];financial_cells=[]
    for s in SCENES:
        xs=sorted(set([0,round(W,3)]+[round(v,3) for b in s.blocks for v in [b.x,b.x+b.w]]+[round(v,3) for o in s.outlines for v in [o.x,o.x+o.w]]))
        ys=sorted(set([0,round(H,3)]+[round(v,3) for b in s.blocks for v in [b.y,b.y+b.h]]+[round(v,3) for o in s.outlines for v in [o.y,o.y+o.h]]))
        grids[s.name]=(xs,ys)
        for b in s.blocks:
            if b.id:
                col=xs.index(round(b.x,3));row=ys.index(round(b.y,3))
                assert b.id not in ids,b.id
                ids[b.id]=(s.name,xl_rowcol_to_cell(row,col,row_abs=True,col_abs=True))
    def resolve(key):
        sheet,address=ids[key]
        return "'"+sheet.replace("'","''")+"'!"+address
    for s in SCENES:
        ws=book.add_worksheet(s.name);xs,ys=grids[s.name]
        ws.hide_gridlines(2);ws.set_paper(9);ws.set_portrait();ws.fit_to_pages(1,1)
        ws.set_margins(0,0,0,0);ws.print_area(0,0,len(ys)-2,len(xs)-2)
        ws.set_zoom(90);ws.set_landscape() if False else None
        ws.set_tab_color(GREEN if s.name=='Company Inventory' else MAROON)
        for i in range(len(xs)-1):
            pixels=(xs[i+1]-xs[i])*96/72
            width=pixels/12 if pixels<12 else (pixels-5)/7
            ws.set_column(i,i,max(.01,width))
        for i in range(len(ys)-1):ws.set_row(i,max(.4,ys[i+1]-ys[i]))
        occupied=[]
        for b in s.blocks:
            c1=xs.index(round(b.x,3));c2=xs.index(round(b.x+b.w,3))-1
            r1=ys.index(round(b.y,3));r2=ys.index(round(b.y+b.h,3))-1
            for a,z,c,d in occupied:
                assert r2<a or r1>z or c2<c or c1>d,(s.name,'overlapping cells',b.text)
            occupied.append((r1,r2,c1,c2))
            props={'font_name':FONT_FACE,'font_size':b.size,'bold':b.bold,'font_color':'#'+b.color,
              'text_wrap':True,'align':b.align,'valign':'top' if not b.border else 'vcenter'}
            if b.fill:props['bg_color']='#'+b.fill
            if b.border:props.update(border=1,border_color='#000000')
            if b.money:props['num_format']='#,##0.00;[Red](#,##0.00);0.00'
            fmt=book.add_format(props)
            if r1!=r2 or c1!=c2:ws.merge_range(r1,c1,r2,c2,'',fmt)
            address=xl_rowcol_to_cell(r1,c1)
            if b.formula:
                expression=b.formula(resolve);ws.write_formula(r1,c1,expression,fmt,'')
                formula_cells.append({'sheet':s.name,'cell':address,'id':b.id,'formula':expression})
            elif b.text!='':
                if isinstance(b.text,(int,float)):ws.write_number(r1,c1,b.text,fmt)
                else:ws.write_string(r1,c1,str(b.text),fmt)
            else:ws.write_blank(r1,c1,None,fmt)
            if b.input:
                ws.write_comment(r1,c1,'Editable input. Unknown company value: left blank by request. Enter a documented value; enter zero only when verified. No application demo data or academic example amount belongs here.')
                ws.data_validation(r1,c1,r2,c2,{'validate':'decimal','criteria':'>=','value':0,'ignore_blank':True,'error_type':'stop','error_message':'Enter a documented non-negative number or leave blank.'})
                input_cells.append({'sheet':s.name,'cell':address,'id':b.id})
            if b.money:financial_cells.append({'sheet':s.name,'cell':address,'id':b.id,'formula':bool(b.formula)})
        # Draw the outer border only into unoccupied edge cells; never overwrite a table.
        for o in s.outlines:
            c1=xs.index(round(o.x,3));c2=xs.index(round(o.x+o.w,3))-1
            r1=ys.index(round(o.y,3));r2=ys.index(round(o.y+o.h,3))-1
            for rr in range(r1,r2+1):
                for cc in range(c1,c2+1):
                    edge={}
                    if rr==r1:edge['top']=2
                    if rr==r2:edge['bottom']=2
                    if cc==c1:edge['left']=2
                    if cc==c2:edge['right']=2
                    if edge and not any(a<=rr<=z and c<=cc<=d for a,z,c,d in occupied):ws.write_blank(rr,cc,None,book.add_format(edge))
        if s.name=='Company Inventory':
            ws.freeze_panes(ys.index(134),xs.index(69))
        elif s.name=='Input Details':ws.freeze_panes(ys.index(171),0)
    book.close()
    manifest={'ids':ids,'formulas':formula_cells,'inputs':input_cells,'financial_cells':financial_cells}
    (VAL/'workbook-map.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    return path,manifest

def powerpoint():
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    for s in SCENES[:4]:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        for o in s.outlines:
            sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(o.x),Pt(o.y),Pt(o.w),Pt(o.h))
            sh.fill.background();sh.line.color.rgb=RGBColor(0,0,0);sh.line.width=Pt(o.width)
        for b in s.blocks:
            if b.fill or b.border:
                sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h))
                if b.id:sh.name=b.id
                if b.fill:sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(b.fill)
                else:sh.fill.background()
                if b.border:sh.line.color.rgb=RGBColor(0,0,0);sh.line.width=Pt(b.border)
                else:sh.line.fill.background()
            if b.text=='':continue
            sh=slide.shapes.add_textbox(Pt(b.x),Pt(b.y),Pt(b.w),Pt(b.h));sh.name=b.id or str(b.text)[:65]
            tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Pt(b.pad)
            for i,line in enumerate(block_lines(b)):
                pp=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                pp.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[b.align]
                pp.space_before=Pt(0);pp.space_after=Pt(0);pp.line_spacing=Pt(b.size*1.12)
                run=pp.add_run();run.text=line;run.font.name=FONT_FACE;run.font.size=Pt(b.size);run.font.bold=b.bold;run.font.color.rgb=RGBColor.from_string(b.color)
        slide.notes_slide.notes_text_frame.text=('Recreated from supplied reference photographs. '+
          'Only the green company monitoring image supplies company data. Financial amounts, development duration and other missing inputs remain blank. '+
          'Application records are demonstration data and are excluded. First three slides reproduce the three reference views; slide four is the green source appendix.')
    prs.core_properties.title='SiteTrace Cost-Benefit Worksheets — Reference Recreation'
    prs.core_properties.subject='Three reference views and green company inventory appendix; financial amounts blank'
    prs.core_properties.author='SiteTrace Academic Documentation'
    prs.core_properties.created=datetime(2026,10,24,tzinfo=timezone.utc)
    path=OUT/'SiteTrace_Cost_Benefit_Worksheets.pptx';prs.save(path);return path

def word_document():
    """Word-native editable text boxes and vector borders at the same page coordinates."""
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import parse_xml
    doc=Document();sec=doc.sections[0]
    sec.page_width=Pt(W);sec.page_height=Pt(H)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Pt(0)
    sec.header_distance=sec.footer_distance=Pt(0)
    doc.styles['Normal'].font.name=FONT_FACE;doc.styles['Normal'].font.size=Pt(9)
    namespace=('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
      'xmlns:w10="urn:schemas-microsoft-com:office:word"')
    serial=0
    for page_index,s in enumerate(SCENES[:4]):
        if page_index:doc.add_page_break()
        anchor=doc.add_paragraph();anchor.paragraph_format.space_before=Pt(0);anchor.paragraph_format.space_after=Pt(0)
        anchor.paragraph_format.line_spacing=Pt(1)
        for o in s.outlines:
            serial+=1
            style=f'position:absolute;margin-left:{o.x}pt;margin-top:{o.y}pt;width:{o.w}pt;height:{o.h}pt;z-index:{serial};mso-position-horizontal-relative:page;mso-position-vertical-relative:page'
            xml=f'<w:pict {namespace}><v:rect id="shape{serial}" style="{style}" filled="f" strokecolor="#000000" strokeweight="{o.width}pt"><w10:wrap type="none"/></v:rect></w:pict>'
            anchor.add_run()._r.append(parse_xml(xml))
        for b in s.blocks:
            serial+=1
            style=f'position:absolute;margin-left:{b.x}pt;margin-top:{b.y}pt;width:{b.w}pt;height:{b.h}pt;z-index:{serial};mso-position-horizontal-relative:page;mso-position-vertical-relative:page;mso-fit-shape-to-text:f'
            attrs=f'filled="{"t" if b.fill else "f"}" stroked="{"t" if b.border else "f"}"'
            if b.fill:attrs+=f' fillcolor="#{b.fill}"'
            if b.border:attrs+=f' strokecolor="#000000" strokeweight="{b.border}pt"'
            paras=[]
            for line in (block_lines(b) or ['']):
                props=(f'<w:spacing w:before="0" w:after="0" w:line="{round(b.size*1.12*20)}" w:lineRule="exact"/>'
                  f'<w:jc w:val="{b.align}"/>')
                run=(f'<w:rPr><w:rFonts w:ascii="{FONT_FACE}" w:hAnsi="{FONT_FACE}"/>'
                  f'<w:sz w:val="{round(b.size*2)}"/><w:color w:val="{b.color}"/>'+('<w:b/>' if b.bold else '')+'</w:rPr>')
                paras.append(f'<w:p><w:pPr>{props}</w:pPr><w:r>{run}<w:t xml:space="preserve">{escape(line)}</w:t></w:r></w:p>')
            content=''.join(paras) or '<w:p/>'
            xml=(f'<w:pict {namespace}><v:rect id="shape{serial}" style="{style}" {attrs}>'
              f'<v:textbox inset="{b.pad}pt,{b.pad}pt,{b.pad}pt,{b.pad}pt"><w:txbxContent>{content}</w:txbxContent></v:textbox><w10:wrap type="none"/></v:rect></w:pict>')
            anchor.add_run()._r.append(parse_xml(xml))
    doc.core_properties.title='SiteTrace Cost-Benefit Worksheets — Reference Recreation'
    doc.core_properties.subject='Only green-table company data; blank financial inputs'
    doc.core_properties.author='SiteTrace Academic Documentation'
    doc.core_properties.created=datetime(2026,10,24,tzinfo=timezone.utc)
    path=OUT/'SiteTrace_Cost_Benefit_Report.docx';doc.save(path);return path

def evaluate_formulas(path):
    import formulas
    from openpyxl.utils import get_column_letter
    result=formulas.ExcelModel().loads(str(path)).finish().calculate()
    values={}
    for rng in result.values():
        if not hasattr(rng,'ranges') or len(rng.ranges)!=1:continue
        spec=rng.ranges[0]
        if 'sheet' not in spec:continue
        for i,row in enumerate(rng.value.tolist()):
            for j,value in enumerate(row):
                if hasattr(value,'item'):value=value.item()
                values[(spec['sheet'].upper(),get_column_letter(spec['n1']+j)+str(int(spec['r1'])+i))]=value
    return values

def validate_deliverables(xlsx_path,manifest,docx_path,pptx_path,pdf_path):
    from lxml import etree
    from pptx import Presentation
    book=load_workbook(xlsx_path,data_only=False)
    cached=load_workbook(xlsx_path,data_only=True)
    ids=manifest['ids']
    def read(key,workbook=cached):
        sheet,address=ids[key];return workbook[sheet][address.replace('$','')].value
    def put(workbook,key,value):
        sheet,address=ids[key];workbook[sheet][address.replace('$','')]=value
    expected=INVENTORY
    actual=[]
    for i,(name,_,__,___) in enumerate(INVENTORY,1):actual.append((name,read('inventory_qty'+str(i)),read('inventory_need'+str(i)),read('inventory_gap'+str(i))))
    checks={
      'all_26_green_inventory_rows_match':actual==expected,
      'recorded_total_758':sum(x[1] for x in actual)==758,
      'reported_gap_total_463':sum(x[3] or 0 for x in actual)==463,
      'three_reported_gaps_match_requirement_minus_recorded':all(n-q==g for _,q,n,g in actual if n is not None),
      '23_unprovided_requirements_remain_blank':sum(x[2] is None for x in actual)==23,
      '23_unprovided_gap_cells_remain_blank':sum(x[3] is None for x in actual)==23,
      'financial_inputs_all_blank':all(book[x['sheet']][x['cell']].value is None for x in manifest['inputs']),
      'financial_display_cells_all_blank':all(cached[x['sheet']][x['cell']].value is None for x in manifest['financial_cells']),
      'five_worksheets':book.sheetnames==[x.name for x in SCENES],
      'no_external_workbook_links':len(book._external_links)==0,
      # XlsxWriter omits attributes equal to their OOXML schema defaults (1).
      'A4_portrait_fit_each_sheet':all(str(s.page_setup.paperSize)=='9' and s.page_setup.orientation=='portrait' and s.sheet_properties.pageSetUpPr.fitToPage and s.page_setup.fitToWidth in (None,1) and s.page_setup.fitToHeight in (None,1) for s in book),
      'no_geometry_overflow':not validate_geometry()}
    # Evaluate every delivered formula, including all duplicate-view links.
    values=evaluate_formulas(xlsx_path)
    bad=[]
    for f in manifest['formulas']:
        key=(f['sheet'].upper(),f['cell']);value=values.get(key,'UNRESOLVED')
        if value not in ['',None]:bad.append([f['sheet'],f['cell'],str(value)])
        for sheet,cell in re.findall(r"'([^']+)'!(\$?[A-Z]+\$?\d+)",f['formula']):
            assert sheet in book.sheetnames,(sheet,cell)
            assert book[sheet][cell.replace('$','')].coordinate==cell.replace('$',''),(sheet,cell)
    checks['all_63_formulas_evaluate_to_blank']=not bad and len(manifest['formulas'])==63
    # Isolated arithmetic fixtures contain no company data and are never shipped.
    scenarios=[]
    with tempfile.TemporaryDirectory(prefix='sitettrace-reference-check-') as temp:
        fixture=load_workbook(xlsx_path)
        put(fixture,'dev1_rate',100);put(fixture,'dev1_months',2)
        test=Path(temp)/'guard.xlsx';fixture.save(test);vals=evaluate_formulas(test)
        def val(key):
            sheet,address=ids[key];return vals[(sheet.upper(),address.replace('$',''))]
        assert val('dev1_subtotal')==200 and val('development_total')==''
        scenarios.append({'name':'partial development inputs','status':'passed','meaning':'Entered rate × duration calculates, but incomplete total remains blank.'})
        for x in manifest['inputs']:fixture[x['sheet']][x['cell']]=0
        test=Path(temp)/'zero.xlsx';fixture.save(test);vals=evaluate_formulas(test)
        assert all(val(k)==0 for k in ['one_time_total','recurring_total','supplies_reduction','labor_reduction','total_reduction'])
        scenarios.append({'name':'explicit zero inputs','status':'passed','meaning':'Documented zero is supported and distinct from an unknown blank.'})
    for path in [xlsx_path,docx_path,pptx_path]:
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for name in z.namelist():
                if name.endswith('.xml'):etree.fromstring(z.read(name))
    prs=Presentation(pptx_path)
    checks['four_editable_powerpoint_slides']=len(prs.slides)==4
    checks['powerpoint_contains_no_flattened_screenshots']=not any(sh.shape_type==13 for slide in prs.slides for sh in slide.shapes)
    with zipfile.ZipFile(docx_path) as z:
        root=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','v':'urn:schemas-microsoft-com:vml'}
        doc_text=' '.join(root.xpath('//w:t/text()',namespaces=ns))
        checks['word_four_fixed_layout_pages']=len(root.xpath('//w:br[@w:type="page"]',namespaces=ns))==3
        checks['word_contains_editable_text_and_vector_shapes']=len(root.xpath('//v:textbox',namespaces=ns))>100 and not any(x.startswith('word/media/') for x in z.namelist())
    ppt_text=' '.join(sh.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame)
    pdf=pymupdf.open(pdf_path);pdf_text=' '.join(pg.get_text() for pg in pdf)
    checks['PDF_four_A4_pages']=len(pdf)==4 and all(abs(pg.rect.width-W)<1 and abs(pg.rect.height-H)<1 for pg in pdf)
    checks['same_green_figures_in_all_document_editions']=all(all(x in txt for x in ['758','463','369','GI PIPES','SWIVEL CLAMP']) for txt in [doc_text,ppt_text,pdf_text])
    excluded=['178,500','20,500','135,000','636,000','576,000','62,585','BuildRight','Casa Buena','Metropolis','San Gabriel','8.0%','3 MONTHS']
    checks['no_academic_example_amounts_or_application_demo_facts']=not any(x in doc_text+ppt_text+pdf_text for x in excluded)
    outside=[]
    for i,pg in enumerate(pdf):
        for block in pg.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line['spans']:
                    x0,y0,x1,y1=span['bbox']
                    if x0<0 or y0<0 or x1>W+.5 or y1>H+.5:outside.append([i+1,span['text']])
    checks['PDF_text_within_page_bounds']=not outside
    assert all(checks.values()),{'checks':checks,'bad_formula_results':bad,'overflow':outside}
    return {'checks':checks,'inventory_categories':26,'recorded_units':758,'reported_gap_units':463,
      'financial_policy':'Leave missing amounts blank; no estimated rates, salaries, prices or savings.',
      'financial_formula_count':63,'isolated_formula_scenarios':scenarios,
      'native_Office_rendering':'Unavailable: COM returned 80070520. XML packages, editable shapes, vector scene geometry and every PDF page are checked. Excel and Office native pagination may differ slightly.',
      'rendered_pages':4,'PowerPoint_slides':4,'Word_pages_by_explicit_page_breaks':4}

def render_previews(pdf_path):
    from PIL import Image,ImageDraw,ImageFont
    doc=pymupdf.open(pdf_path)
    names=['01-one-time-cost','02-cost-explanation','03-tangible-benefits','04-green-company-inventory']
    contact=Image.new('RGB',(1260,485),'#E8EBEF')
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',12)
    for i,pg in enumerate(doc):
        path=VIEWS/(names[i]+'.png');pg.get_pixmap(matrix=pymupdf.Matrix(2,2)).save(path)
        thumb=Image.open(path).convert('RGB');thumb.thumbnail((300,425))
        contact.paste(thumb,(i*315+7,10));ImageDraw.Draw(contact).text((i*315+8,450),names[i],font=font,fill='black')
    contact.save(VAL/'reference-recreation-contact.png')
    return [str(VIEWS/(n+'.png')) for n in names]

def write_sources():
    source=OUT/'source'
    with (source/'green-company-inventory.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['ITEM NO.','POWER TOOLS','QTY.','NEED TO COMPLY','MISSING OR NO REPORT'])
        for i,items in enumerate(INVENTORY,1):writer.writerow([i,*items])
    originals=[
      ('company data only','dc83a570-9cc7-472d-be16-137f87d2c5e9.jpg'),
      ('reference view 1: cost worksheet layout only','4402c8cb-1525-479f-ba0b-8aa5a8a23ba7.jpg'),
      ('reference view 2: repeated lower worksheet and explanatory paragraph layout only','006f9d77-0cab-47a5-b65f-bcf640b8e2df.jpg'),
      ('reference view 3: tangible-benefits layout only','7a1bd893-e0fc-4dba-ac70-24ada0c81048.jpg')]
    records=[]
    for role,name in originals:
        path=Path('C:/Users/Roland Carl/Downloads')/name
        records.append({'attachment':name,'role':role,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    (VAL/'company-data-source.json').write_text(json.dumps({'policy':'Only green table values are company data. Application data is demonstration-only and excluded. Reference photographs supply layout, not financial values.','sources':records},indent=2),encoding='utf-8')
    (source/'requirements.txt').write_text('XlsxWriter==3.3.0\nopenpyxl==3.1.5\npython-pptx==1.0.2\npython-docx==1.2.0\nreportlab==5.1.0\nPyMuPDF==2.0\nPillow==13.0.0\nformulas==1.3.4\n',encoding='utf-8')

def write_readme():
    text='''# SiteTrace cost-benefit worksheets — reference recreation

This version follows the user's corrected instructions. **Only the green “HAND TOOLS AND OTHERS MONITORING” image supplies company data. All application records are demonstration data. All missing financial amounts remain blank.**

## Files

- `SiteTrace_Cost_Benefit_Worksheets.pptx`: editable PowerPoint text and vector shapes; three reference views plus the green company inventory appendix.
- `SiteTrace_Cost_Benefit_Analysis.xlsx`: editable Excel worksheets with blank financial inputs and 63 guarded formulas.
- `SiteTrace_Cost_Benefit_Report.docx`: editable Word text boxes and vector borders using the same fixed page layout.
- `SiteTrace_Cost_Benefit_Report.pdf`: four A4 pages generated from the same layout.
- `recreated-views/`: clean high-resolution PNG copies of the three recreated views and the green inventory appendix.

## How the photos were recreated

1. **Photo 1:** the full, bordered cost worksheet, with development/training/hardware tables on the left and one-time summary/recurring costs on the right.
2. **Photo 2:** the repeated lower part of the same worksheet, followed by its explanatory paragraph. In Excel these repeated monetary cells link to the first sheet instead of becoming duplicate inputs.
3. **Photo 3:** side-by-side Existing System and Proposed System supplies and labour tables, with the original-style pale blue reduction bands and a paragraph underneath.
4. **Appendix:** the actual green inventory table, including all 26 populated categories and the original blank row 27.

Monitor borders, browser controls, perspective distortion and the Windows taskbar are excluded from the clean document recreation. Maroon chapter bands, plain borders, grey headings, narrow type, two-column layouts and blue result bands follow the reference photographs. Worksheet labels are generic categories, not assertions about actual company employees or purchases. Cashier/CIO examples were replaced with unquantified inventory-related personnel groups. The hardware header uses “Unit Cost” to match its calculation. No financial example amount, three-month duration, training-day quantity, employee count or discount rate from another project is copied.

## Company data

The green table contains 26 populated categories and 758 recorded units. Three categories have 463 missing or unreported units: 19 round shovels, 75 GI pipes and 369 swivel clamps. Each gap equals its supplied requirement minus its recorded quantity. Blank requirements and gaps stay blank. The source spelling “CATWALK FLATFORM SCAFFOLDING” is preserved. Missing/no-report units are reconciliation gaps, not confirmed physical losses or monetary losses.

## Excel editing

The first three sheets reproduce the reference views. `Company Inventory` contains the supplied data. `Input Details` holds the unknown training-day and hardware/software quantities required by the formulas. Every financial input has an editable-cell comment and numeric validation. It is intentionally white to preserve the reference appearance.

- Enter documented rates, durations, quantities and monthly labour totals only when available.
- Training Kit is a total materials charge; no kit count is assumed.
- Hardware subtotals require a unit cost on the cost sheet and quantity on Input Details.
- Monthly labour fields represent the combined monthly cost of each named personnel group; no employee count is assumed. Annualization multiplies by 12 as a unit conversion.
- Maintenance and communication rates are monthly and license renewal is annual, following the reference's billing units; these are not claims about an actual subscription.
- All subtotals, totals, differences and repeated-view values remain blank until their necessary inputs are numeric. A documented zero is supported and is distinct from an unknown blank.
- Payroll reduction cannot be inferred from faster reporting. No headcount reduction, equipment price, loss recovery or annual saving is assumed.

PowerPoint, Word and PDF are static blank versions of the worksheets. They do not live-link to later Excel edits. Their text and borders can be edited in PowerPoint/Word; update those editions alongside a completed Excel scenario.

## Validation and previous version

`validation/validation-report.json` records quantity checks, missing-input guards, all formula results, cross-sheet references, isolated arithmetic checks, XML package integrity, editable document content, page geometry and source exclusions. `validation/company-data-source.json` contains only the four supplied image sources and hashes. No application dataset or repository documentation is a company-data source.

All pages are inspected through the PDF and common vector layout. Native Office automation was unavailable in this Windows session (COM 80070520); the editable Office packages are checked structurally. Word/PowerPoint use native vector shapes and editable text, not pasted page screenshots. Minor native font/pagination differences may occur. Excel uses an A4 portrait grid fitted to one printed page per sheet.

The superseded initial analysis is retained in `archive/initial-analysis/`. It is not an input to this recreation. No application source code, database schema or production configuration was modified.

Rebuild with Python 3.11 and `source/requirements.txt` by running `python source/build_reference_recreation.py`. This restores the requested blank baseline; back up any later workbook edits first. The original four attachments are expected at their supplied Downloads paths only for provenance hashing.
'''
    (OUT/'README.md').write_text(text,encoding='utf-8')

def main():
    for p in [OUT,VAL,VIEWS]:p.mkdir(parents=True,exist_ok=True)
    fit_padding();assert not validate_geometry(),validate_geometry()
    write_sources()
    xlsx,manifest=excel_workbook();print('Reference-layout Excel workbook built',flush=True)
    pdf=OUT/'SiteTrace_Cost_Benefit_Report.pdf';render_pdf(pdf,SCENES[:4])
    pptx=powerpoint();docx=word_document()
    print('PowerPoint, Word and PDF recreated',flush=True)
    report=validate_deliverables(xlsx,manifest,docx,pptx,pdf)
    report['previews']=render_previews(pdf)
    report['artifacts']=[{'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [xlsx,docx,pdf,pptx]]
    report['user_decision']='Leave missing amounts blank.'
    report['source_scope']='Only the supplied green inventory table supplies company quantities. No system demo data, salaries, prices, dates, schedules, project names or procurement figures were imported.'
    (VAL/'validation-report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    write_readme()
    print(json.dumps({'pages':4,'slides':4,'worksheets':5,'formulas':63,'company_units':758,'gap_units':463,'financial_inputs':'blank','files':[str(p) for p in [xlsx,pptx,docx,pdf]]},indent=2),flush=True)

if __name__=='__main__':main()
